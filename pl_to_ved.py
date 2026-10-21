#!/usr/bin/env python3
"""
Перенос данных из упаковочного листа (Packing List, .xlsx) в таблицу
для «Формы загрузки ВЭД» в 1С.

Результат — новый Excel-файл, одна строка на контейнер, без пустых строк:

    Контейнер | Номер пломбы | Вес | Количество | Страна

Вес и количество по контейнеру суммируются по всем товарным строкам
контейнера и сверяются со строками SUBTOTAL / TOTAL упаковочного листа.
Страна определяется по строкам «COUNTRY OF ORIGIN», «MANUFACTURER»
или «PORT OF LOADING» (её можно задать вручную ключом --country).

Использование:
    python pl_to_ved.py PSS-VN-2607-7_PL.xlsx
    python pl_to_ved.py PL1.xlsx PL2.xlsx
    python pl_to_ved.py PL.xlsx -o результат.xlsx --country Вьетнам
    python pl_to_ved.py            (откроется окно выбора файлов)
"""

import argparse
import re
import string
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_HEADERS = ["Контейнер", "Номер пломбы", "Вес", "Количество", "Страна"]
OUTPUT_SUFFIX = "_ВЭД"

# Английское название страны в документе -> название для 1С.
# Более длинные названия стоят раньше, чтобы «HONG KONG» не превратился в «CHINA» и т.п.
COUNTRIES = [
    ("HONG KONG", "Гонконг"),
    ("HONGKONG", "Гонконг"),
    ("SOUTH KOREA", "Корея, Республика"),
    ("KOREA", "Корея, Республика"),
    ("VIET NAM", "Вьетнам"),
    ("VIETNAM", "Вьетнам"),
    ("CHINA", "Китай"),
    ("TAIWAN", "Тайвань (Китай)"),
    ("THAILAND", "Таиланд"),
    ("MALAYSIA", "Малайзия"),
    ("INDONESIA", "Индонезия"),
    ("CAMBODIA", "Камбоджа"),
    ("INDIA", "Индия"),
    ("JAPAN", "Япония"),
    ("TURKEY", "Турция"),
    ("TURKIYE", "Турция"),
    ("SERBIA", "Сербия"),
    ("BELARUS", "Беларусь"),
    ("KAZAKHSTAN", "Казахстан"),
    ("UZBEKISTAN", "Узбекистан"),
]

# Подписи в упаковочном листе, рядом с которыми ищем страну (по убыванию приоритета).
COUNTRY_LABELS = ["COUNTRY OF ORIGIN", "MADE IN", "MANUFACTURER", "PORT OF LOADING"]

# Номер контейнера ISO 6346: 3 буквы владельца + U/J/Z + 6 цифр + контрольная цифра.
CONTAINER_RE = re.compile(r"[A-Z]{3}[UJZ]\s?\d{6}\s?\d")

WEIGHT_HEADER_KEYS = ["G.W", "GROSS", "N.W", "NET", "WEIGHT"]


class PackingListError(Exception):
    pass


@dataclass
class Container:
    number: str
    seal: str
    row: int
    quantity: float = 0
    weight: float = 0
    subtotal_quantity: float = None
    subtotal_weight: float = None


def _text(value):
    if value is None:
        return ""
    return " ".join(str(value).upper().split())


def to_number(value):
    """Число из ячейки: 15749, '15749', '15 749', '15,749.5', '15749/15000' (берётся первое)."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).split("/")[0].strip().replace("\xa0", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+", s):
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _clean_number(x):
    x = round(x, 3)
    return int(x) if x == int(x) else x


def _container_check_digit_ok(number):
    values = {}
    n = 10
    for ch in string.ascii_uppercase:
        if n % 11 == 0:
            n += 1
        values[ch] = n
        n += 1
    values.update({d: int(d) for d in string.digits})
    total = sum(values[ch] * 2 ** i for i, ch in enumerate(number[:10]))
    return total % 11 % 10 == int(number[10])


def split_container_and_seal(value):
    """'CICU7418771/XHL200914' -> ('CICU7418771', 'XHL200914')."""
    raw = str(value).strip()
    m = CONTAINER_RE.search(raw.upper())
    if m:
        number = m.group().replace(" ", "")
        rest = raw[:m.start()] + " " + raw[m.end():]
    else:
        parts = re.split(r"[/\n]", raw, maxsplit=1)
        number = parts[0].strip()
        rest = parts[1] if len(parts) > 1 else ""
    seal = re.sub(r"\s*[\n/]\s*", "/", rest.strip(" /\\\n\t,;:-"))
    return number, seal


def find_table(ws):
    """Ищет строку заголовка таблицы. Возвращает (строка, кол. контейнера, кол. количества, кол. веса)."""
    for row in ws.iter_rows(max_row=min(ws.max_row, 100)):
        container_col = qty_col = None
        weight_cols = {}
        for cell in row:
            t = _text(cell.value)
            if not t:
                continue
            if container_col is None and ("CONTAINER" in t or "КОНТЕЙНЕР" in t):
                container_col = cell.column
            elif qty_col is None and ("QUANTITY" in t or t.startswith("QTY") or "КОЛ" in t):
                qty_col = cell.column
            else:
                for priority, key in enumerate(WEIGHT_HEADER_KEYS):
                    if key in t:
                        weight_cols.setdefault(priority, cell.column)
                        break
        if container_col and qty_col and weight_cols:
            return row[0].row, container_col, qty_col, weight_cols[min(weight_cols)]
    return None


def find_country(ws):
    """Ищет страну рядом с подписями COUNTRY_LABELS. Возвращает (страна, подпись) или (None, None)."""
    rows = []
    for row in ws.iter_rows():
        text = " ".join(_text(c.value) for c in row if c.value not in (None, ""))
        if text:
            rows.append(text)
    for label in COUNTRY_LABELS:
        for text in rows:
            pos = text.find(label)
            if pos < 0:
                continue
            tail = text[pos + len(label):]
            for en, ru in COUNTRIES:
                if re.search(r"\b" + re.escape(en) + r"\b", tail):
                    return ru, label
    return None, None


def parse_packing_list(path):
    """Читает упаковочный лист. Возвращает (контейнеры, страна, источник страны, TOTAL, предупреждения)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        table = find_table(ws)
        if table:
            break
    else:
        raise PackingListError(
            "не найдена таблица с колонками Container / Quantity / G.W (вес)"
        )
    header_row, container_col, qty_col, weight_col = table

    warnings = []
    containers = []
    current = None
    total = None
    for row in ws.iter_rows(min_row=header_row + 1):
        values = {c.column: c.value for c in row}
        row_no = row[0].row
        compact = "".join(_text(v) for v in values.values()).replace(" ", "")
        qty = to_number(values.get(qty_col))
        weight = to_number(values.get(weight_col))

        if "SUBTOTAL" in compact:
            if current is not None:
                current.subtotal_quantity, current.subtotal_weight = qty, weight
            continue
        if "TOTAL" in compact:
            total = (qty, weight)
            break

        if values.get(container_col) not in (None, ""):
            number, seal = split_container_and_seal(values[container_col])
            current = Container(number, seal, row_no)
            containers.append(current)
        if qty is None and weight is None:
            continue
        if current is None:
            warnings.append(f"строка {row_no}: товар без номера контейнера — пропущен")
            continue
        current.quantity += qty or 0
        current.weight += weight or 0

    if not containers:
        raise PackingListError("в таблице не найдено ни одного контейнера")

    for c in containers:
        if not CONTAINER_RE.fullmatch(c.number):
            warnings.append(f"строка {c.row}: «{c.number}» не похож на номер контейнера")
        elif not _container_check_digit_ok(c.number):
            warnings.append(f"строка {c.row}: у контейнера {c.number} неверная контрольная цифра")
        if not c.seal:
            warnings.append(f"строка {c.row}: у контейнера {c.number} нет номера пломбы")
        if c.subtotal_quantity is not None and round(c.subtotal_quantity, 3) != round(c.quantity, 3):
            warnings.append(
                f"{c.number}: сумма количества {_clean_number(c.quantity)} "
                f"≠ SUBTOTAL {_clean_number(c.subtotal_quantity)}"
            )
        if c.subtotal_weight is not None and round(c.subtotal_weight, 3) != round(c.weight, 3):
            warnings.append(
                f"{c.number}: сумма веса {_clean_number(c.weight)} "
                f"≠ SUBTOTAL {_clean_number(c.subtotal_weight)}"
            )

    country, country_source = find_country(ws)
    return containers, country, country_source, total, warnings


def write_result(path, containers, country):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ВЭД"
    ws.append(OUTPUT_HEADERS)
    for c in containers:
        ws.append([c.number, c.seal, _clean_number(c.weight), _clean_number(c.quantity), country])

    header_fill = PatternFill("solid", fgColor="FFD200")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for letter, width in zip("ABCDE", (18, 18, 12, 12, 14)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


def convert(src, dst=None, country=None):
    src = Path(src)
    if src.suffix.lower() == ".xls":
        raise PackingListError("старый формат .xls не поддерживается — пересохраните файл как .xlsx")
    if dst is None:
        dst = src.with_name(src.stem + OUTPUT_SUFFIX + ".xlsx")

    containers, found_country, country_source, total, warnings = parse_packing_list(src)
    if country is None:
        country = found_country or ""
        if not found_country:
            warnings.append("страна не найдена — колонка «Страна» пустая (задайте её ключом --country)")

    sum_qty = sum(c.quantity for c in containers)
    sum_weight = sum(c.weight for c in containers)
    if total is not None:
        total_qty, total_weight = total
        if total_qty is not None and round(total_qty, 3) != round(sum_qty, 3):
            warnings.append(f"итого количество {_clean_number(sum_qty)} ≠ TOTAL {_clean_number(total_qty)}")
        if total_weight is not None and round(total_weight, 3) != round(sum_weight, 3):
            warnings.append(f"итого вес {_clean_number(sum_weight)} ≠ TOTAL {_clean_number(total_weight)}")

    try:
        write_result(dst, containers, country)
    except PermissionError:
        raise PackingListError(f"не удалось записать {dst} — закройте файл в Excel и повторите")

    print(f"\n{src.name} -> {dst}")
    print(f"{'Контейнер':<14}{'Пломба':<16}{'Вес':>10}{'Кол-во':>9}  Страна")
    for c in containers:
        print(f"{c.number:<14}{c.seal:<16}{_clean_number(c.weight):>10}"
              f"{_clean_number(c.quantity):>9}  {country}")
    print(f"Итого: контейнеров {len(containers)}, вес {_clean_number(sum_weight)}, "
          f"количество {_clean_number(sum_qty)}")
    if country and country_source:
        print(f"Страна «{country}» определена по строке {country_source}")
    for w in warnings:
        print(f"ВНИМАНИЕ: {w}")
    return dst


def ask_files():
    try:
        import tkinter
        from tkinter import filedialog
    except ImportError:
        return []
    root = tkinter.Tk()
    root.withdraw()
    files = filedialog.askopenfilenames(
        title="Выберите упаковочный лист (Packing List)",
        filetypes=[("Excel", "*.xlsx *.xlsm"), ("Все файлы", "*.*")],
    )
    root.destroy()
    return [Path(f) for f in files]


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Перенос контейнеров, пломб, веса, количества и страны "
                    "из Packing List в Excel для формы загрузки ВЭД 1С."
    )
    parser.add_argument("files", nargs="*", type=Path, help="упаковочные листы .xlsx")
    parser.add_argument("-o", "--output", type=Path,
                        help="имя файла результата (только для одного входного файла)")
    parser.add_argument("--country", help="страна вручную, например «Вьетнам»")
    args = parser.parse_args(argv)

    interactive = not args.files
    files = ask_files() if interactive else args.files
    if not files:
        parser.print_help()
        return 1
    if args.output and len(files) > 1:
        parser.error("--output можно указать только для одного входного файла")

    rc = 0
    for f in files:
        try:
            convert(f, args.output, args.country)
        except (PackingListError, OSError) as e:
            print(f"\nОШИБКА ({f}): {e}")
            rc = 1
    if interactive:
        input("\nНажмите Enter, чтобы закрыть окно...")
    return rc


if __name__ == "__main__":
    sys.exit(main())
